# encoding: utf-8
"""
BIM Health Scorecard - Excel writer (host-side, run with system Python + openpyxl).
Produces a formatted .xlsx scorecard: Summary, Rule Detail, Failing Elements.

Data source (either):
  --fetch   ask the AutoConst Revit MCP route /health_audit/ directly (uncapped,
            no AI tokens), save it to outputs/scorecard_data.json, then build.
  default   read outputs/scorecard_data.json (AutoConst tool JSON or the legacy
            v1.3 script JSON - both formats are accepted).

Usage:  py build_scorecard.py --fetch [--stage design|fabrication|operations]
        py build_scorecard.py
"""
import argparse, json, os, sys
from urllib import request as urlrequest
from urllib.error import URLError, HTTPError

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "outputs", "scorecard_data.json")

GRADE_FILL = {  # letter -> hex fill
    "A": "2E7D32", "B": "7CB342", "C": "F9A825", "D": "EF6C00", "F": "C62828",
}
HDR_FILL = "1F3A5F"
WHITE = "FFFFFF"
GREY = "F2F2F2"
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def grade_fill(letter):
    return PatternFill("solid", fgColor=GRADE_FILL.get(letter, "9E9E9E"))

def passfill(pct):
    if pct >= 90: c = "C8E6C9"
    elif pct >= 70: c = "FFF9C4"
    elif pct >= 40: c = "FFE0B2"
    else: c = "FFCDD2"
    return PatternFill("solid", fgColor=c)

def hdr(cell):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=HDR_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

API = os.environ.get("REVIT_MCP_API", "http://localhost:48884/revit_mcp")
CAT_WEIGHTS = {"naming": 0.15, "parameters": 0.35, "classification": 0.20,
               "geometry": 0.20, "worksharing": 0.10}
GRADES = [("A", 90), ("B", 80), ("C", 70), ("D", 60), ("F", 0)]


def letter(score):
    for g, lo in GRADES:
        if score >= lo:
            return g
    return "F"


def fetch(stage):
    """Uncapped audit straight from the Revit route; returns the parsed JSON."""
    body = json.dumps({"stage": stage, "max_failing_ids": 0}).encode("utf-8")
    req = urlrequest.Request(API + "/health_audit/", data=body,
                             headers={"Content-Type": "application/json"})
    try:
        with urlrequest.urlopen(req, timeout=300) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except HTTPError as e:
        sys.exit("Audit failed (HTTP %d): %s" % (e.code, e.read().decode("utf-8", "replace")[:500]))
    except URLError as e:
        sys.exit("Cannot reach Revit at %s (%s). Is Revit open with pyRevit Routes running?"
                 % (API, e.reason))


def normalize(raw):
    """Map AutoConst tool JSON (or legacy v1.3 JSON) to the shape this builder writes.

    Returns (d, sample_note). sample_note is None when the failing-elements list is
    complete; otherwise a description of what was left out.
    """
    if "engine_version" not in raw:  # legacy v1.3 script output - already in builder shape
        return raw, None
    if raw.get("status") != "success":
        sys.exit("Audit returned an error: %s" % raw.get("error", raw))

    cats = {}
    for cat, score in raw["category_scores"].items():
        cats[cat] = {"score": score, "grade": letter(score), "weight": CAT_WEIGHTS.get(cat, 0)}

    rules, omitted = [], []
    for r in raw["rules"]:
        if not r.get("applicable", True):
            continue  # category has no elements in this model (v1.3 skipped these too)
        rules.append({
            "cat": r["category"], "id": r["id"], "name": r["name"], "src": r["source"],
            "sev": r["severity"], "na": r["advisory"], "passrate": r["pass_rate"],
            "fail_count": r["fail_count"],
            "fails": [(f["id"], f["reason"]) for f in r["failing"]],
        })
        if r["failing_truncated"]:
            omitted.append("%s: %d of %d listed" % (r["id"], r["failing_returned"],
                                                    r["findings_available"]))

    d = {"model": raw["document"]["title"], "stage": raw["stage"], "overall": raw["overall"],
         "grade": raw["grade"], "category_scores": cats, "rules": rules}

    if not raw.get("complete_export", True) or omitted:
        note = "SAMPLE - capped. Not every failing element is listed. " + "; ".join(omitted)
        if raw.get("error_code"):
            note = "%s (%s)" % (note, raw["error_code"])
        return d, note

    listed = sum(len(r["fails"]) for r in rules)
    if listed != raw["total_findings"]:
        sys.exit("Failing-element list incomplete: %d rows vs total_findings %d"
                 % (listed, raw["total_findings"]))
    return d, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fetch", action="store_true",
                    help="run the audit via the Revit route and save it first")
    ap.add_argument("--stage", default="design", choices=["design", "fabrication", "operations"])
    args = ap.parse_args()

    if args.fetch:
        raw = fetch(args.stage)
        os.makedirs(os.path.dirname(DATA), exist_ok=True)
        with open(DATA, "w", encoding="utf-8") as f:
            json.dump(raw, f, indent=1, ensure_ascii=False)
    else:
        with open(DATA, encoding="utf-8") as f:
            raw = json.load(f)
    d, sample_note = normalize(raw)

    wb = Workbook()

    # ---------- Sheet 1: Summary ----------
    ws = wb.active
    ws.title = "Scorecard"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col, w in zip("BCDEF", (26, 14, 12, 12, 30)):
        ws.column_dimensions[col].width = w

    ws["B2"] = "BIM HEALTH SCORECARD"
    ws["B2"].font = Font(bold=True, size=20, color=HDR_FILL)
    ws["B3"] = d["model"]
    ws["B3"].font = Font(bold=True, size=12, color="555555")
    ws["B4"] = "Stage: %s   |   generated by AutoConst" % d["stage"]
    ws["B4"].font = Font(size=10, italic=True, color="888888")

    # big overall grade block
    ws.merge_cells("B6:C8")
    g = ws["B6"]
    g.value = "%s\n%.1f / 100" % (d["grade"], d["overall"])
    g.font = Font(bold=True, size=28, color=WHITE)
    g.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(6, 9):
        for c in ("B", "C"):
            ws["%s%d" % (c, r)].fill = grade_fill(d["grade"])
    ws["D6"] = "OVERALL"
    ws["D6"].font = Font(bold=True, size=12, color=HDR_FILL)

    # category table
    ws["B10"] = "Category"; ws["C10"] = "Score"; ws["D10"] = "Grade"; ws["E10"] = "Weight"
    for c in "BCDE":
        hdr(ws["%s10" % c])
    order = ["naming", "parameters", "classification", "geometry", "worksharing"]
    row = 11
    cats = d["category_scores"]
    for cat in order:
        if cat not in cats:
            continue
        cs = cats[cat]
        ws["B%d" % row] = cat.capitalize()
        ws["C%d" % row] = cs["score"]
        ws["D%d" % row] = cs["grade"]
        ws["E%d" % row] = "%d%%" % round(cs["weight"] * 100)
        ws["C%d" % row].fill = passfill(cs["score"])
        ws["D%d" % row].fill = grade_fill(cs["grade"])
        ws["D%d" % row].font = Font(bold=True, color=WHITE)
        for c in "BCDE":
            cell = ws["%s%d" % (c, row)]
            cell.border = BORDER
            if c != "B":
                cell.alignment = Alignment(horizontal="center")
        row += 1

    note = ("Grades: A>=90 B>=80 C>=70 D>=60 F<60. Weights are AutoConst methodology "
            "(tunable per BIM Execution Plan). Every rule cites its source on the "
            "Rule Detail tab. Advisory (info) rules are reported but not scored.")
    ws.merge_cells("B%d:F%d" % (row + 1, row + 3))
    nc = ws["B%d" % (row + 1)]
    nc.value = note
    nc.font = Font(size=9, italic=True, color="777777")
    nc.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- Sheet 2: Rule Detail ----------
    rd = wb.create_sheet("Rule Detail")
    rd.sheet_view.showGridLines = False
    cols = ["Category", "Rule", "Check", "Source", "Severity", "Pass %", "# Fails"]
    widths = [14, 12, 34, 14, 10, 9, 9]
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = rd.cell(row=1, column=i, value=c); hdr(cell)
        rd.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for rule in d["rules"]:
        pct = round(rule["passrate"] * 100)
        sev = rule["sev"] + (" (adv)" if rule.get("na") else "")
        vals = [rule["cat"], rule["id"], rule["name"], rule["src"], sev, pct,
                rule.get("fail_count", len(rule["fails"]))]
        for i, v in enumerate(vals, start=1):
            cell = rd.cell(row=r, column=i, value=v)
            cell.border = BORDER
            if i in (5, 6, 7):
                cell.alignment = Alignment(horizontal="center")
        rd.cell(row=r, column=6).fill = passfill(pct)
        if r % 2 == 0:
            for i in range(1, 8):
                if i != 6:
                    rd.cell(row=r, column=i).fill = PatternFill("solid", fgColor=GREY)
        r += 1
    rd.freeze_panes = "A2"

    # ---------- Sheet 3: Failing Elements ----------
    fe = wb.create_sheet("Failing Elements (SAMPLE)" if sample_note else "Failing Elements")
    fe.sheet_view.showGridLines = False
    fcols = ["Rule", "Check", "Element Id", "Issue"]
    fwidths = [12, 34, 14, 40]
    first = 1
    if sample_note:
        fe.merge_cells("A1:D1")
        fe["A1"] = sample_note
        fe["A1"].font = Font(bold=True, color="C62828")
        fe["A1"].alignment = Alignment(wrap_text=True)
        fe.row_dimensions[1].height = 45
        first = 2
    for i, (c, w) in enumerate(zip(fcols, fwidths), start=1):
        cell = fe.cell(row=first, column=i, value=c); hdr(cell)
        fe.column_dimensions[get_column_letter(i)].width = w
    r = first + 1
    for rule in d["rules"]:
        for eid, issue in rule["fails"]:
            fe.cell(row=r, column=1, value=rule["id"]).border = BORDER
            fe.cell(row=r, column=2, value=rule["name"]).border = BORDER
            fe.cell(row=r, column=3, value=eid).border = BORDER
            fe.cell(row=r, column=4, value=issue).border = BORDER
            r += 1
    if r == first + 1:
        fe.cell(row=r, column=1, value="No failing elements recorded.")
    fe.freeze_panes = "A%d" % (first + 1)

    model_safe = "".join(ch for ch in d["model"] if ch.isalnum() or ch in " -_").strip().replace(" ", "-")
    out = os.path.join(ROOT, "outputs", "BIM-Health-Scorecard-%s.xlsx" % model_safe)
    wb.save(out)
    print("WROTE", out)
    print("GRADE %s  %.1f/100  stage=%s  failing rows=%d%s" % (
        d["grade"], d["overall"], d["stage"], r - first - 1,
        "  (SAMPLE - capped)" if sample_note else "  (complete)"))

if __name__ == "__main__":
    main()
